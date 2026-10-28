"""Build the final BCSE497J Project-I report from the supplied VIT template.

    python scripts/build_final_report.py            # build, then fill the TOC page numbers

The template's cover page, declaration, certificate, acknowledgements, list
tables, abstract and chapter shells are filled in place so that VIT's own
formatting, headers and footers are kept. Chapter text lives in
final_report_text.py. Page numbers in the contents table are read back from a
Word-rendered PDF (Word must be installed), so they match the real layout.
"""
from __future__ import annotations

import copy
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
import final_report_text as T  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "BCSE497J Project I Report - Template FINAlL Updated  (2) (1).docx"
OUTPUT = ROOT / "Agent SecOps - BCSE497J Project-I Final Report (Humanised).docx"
FIG = ROOT / "docs" / "figures"
SHOT = ROOT / "docs" / "screenshots"
CROP = SHOT / "report"
FONT = "Times New Roman"


# ----------------------------------------------------------------------
# images
# ----------------------------------------------------------------------
def _crop(src: str, dst: str, box):
    CROP.mkdir(parents=True, exist_ok=True)
    im = Image.open(SHOT / src).convert("RGB")
    w, h = im.size
    l, t, r, b = box
    im.crop((l, t, r, min(b, h))).save(CROP / dst)
    return CROP / dst


def image_paths():
    return {
        "gantt": FIG / "fig1_gantt.png",
        "arch": FIG / "fig2_architecture.png",
        "dfd": FIG / "fig4_dfd.png",
        "usecase": FIG / "fig5_usecase.png",
        "class": FIG / "fig6_class.png",
        "sequence": FIG / "fig7_sequence.png",
        "bench": FIG / "fig8_benchmark.png",
        "shot_pass": _crop("dash_compliant_pass.png", "pass.png", (0, 700, 1640, 2500)),
        "shot_team_top": _crop("dash_team_demo.png", "team_top.png", (0, 700, 1640, 2230)),
        "shot_team_findings": _crop("dash_team_demo.png", "team_findings.png", (0, 2920, 1640, 4700)),
        "shot_off": _crop("dash_reasoning_off.png", "off.png", (0, 700, 1640, 2100)),
        "shot_tf": _crop("dash_terraform_aws.png", "tf.png", (0, 700, 1640, 2840)),
        "shot_bench": _crop("dash_benchmark.png", "bench.png", (0, 650, 1640, 2533)),
        "gh_run": SHOT / "github_actions_run7.png",
        "gh_approve": SHOT / "github_approve_dialog.png",
        "azure_rg": SHOT / "azure_resource_groups.png",
    }


# ----------------------------------------------------------------------
# low-level formatting helpers
# ----------------------------------------------------------------------
def set_font(run, size=None, bold=None, italic=None, name=FONT, color=None):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        set_font(p.add_run(text), 12)
        return p
    first = runs[0]
    first.text = text
    first.font.highlight_color = None
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return p


def remove(p):
    el = p._p if hasattr(p, "_p") else p
    el.getparent().remove(el)


def clone_after(p, text):
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    set_text(np_, text)
    return np_


def spacing(p, line=1.15, before=0, after=6, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "808080")
        b.append(e)
    tblpr.append(b)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trpr.append(c)


# ----------------------------------------------------------------------
# block builders: create at the document end, then move to the anchor
# ----------------------------------------------------------------------
class Writer:
    def __init__(self, doc, anchor_el, images):
        self.doc = doc
        self.anchor = anchor_el
        self.images = images
        self.headings = []   # (level, label, text)  for the TOC pass

    def _place(self, el):
        self.anchor.addprevious(el)

    def para(self, text="", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             line=1.15, before=0, after=6, keep_next=False, page_break=False):
        p = self.doc.add_paragraph()
        if text:
            r = p.add_run(text)
            set_font(r, size, bold, italic)
        spacing(p, line, before, after, align, keep_next)
        if page_break:
            p.paragraph_format.page_break_before = True
        self._place(p._p)
        return p

    def chapter(self, label, title):
        self.para(label, 16, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, after=4,
                  keep_next=True, page_break=True)
        num = label.split()[-1]
        self.para("{}. {}".format(num, title), 14, True, align=WD_ALIGN_PARAGRAPH.LEFT,
                  line=1.5, after=8, keep_next=True)
        self.headings.append((1, "{}.".format(num), title))

    def h2(self, text):
        self.para(text, 13, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, before=6, after=4,
                  keep_next=True)
        self.headings.append((2, text.split(" ", 1)[0], text.split(" ", 1)[1]))

    def h3(self, text):
        self.para(text, 12, True, True, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, before=4,
                  after=3, keep_next=True)
        head = text.split(" ", 1)
        if head[0][0].isdigit():
            self.headings.append((3, head[0], head[1]))

    def bullets(self, items, numbered=False):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            marker = "{}.".format(i) if numbered else "•"
            r = p.add_run("{}\t{}".format(marker, item))
            set_font(r, 12)
            spacing(p, 1.15, 0, 4, WD_ALIGN_PARAGRAPH.JUSTIFY)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.4)
            pf.first_line_indent = Inches(-0.3)
            pf.tab_stops.add_tab_stop(Inches(0.4))
            self._place(p._p)

    def table(self, caption, header, rows, widths):
        self.para(caption, 10, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, before=6,
                  after=3, keep_next=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        for ci, h in enumerate(header):
            c = t.rows[0].cells[ci]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            set_font(r, 10.5, True)
            shade(c, "D9D9D9")
        repeat_header(t.rows[0])
        for ri, row in enumerate(rows, 1):
            no_split(t.rows[ri])
            for ci, val in enumerate(row):
                c = t.rows[ri].cells[ci]
                c.text = ""
                r = c.paragraphs[0].add_run(str(val))
                set_font(r, 10.5, str(row[0]).startswith("Total"))
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Inches(w)
            for c in row.cells:
                for p in c.paragraphs:
                    spacing(p, 1.0, 1, 1, WD_ALIGN_PARAGRAPH.LEFT)
        self._place(t._tbl)
        self.para("", 6, after=4)

    def figure(self, key, caption):
        path = self.images[key]
        with Image.open(path) as im:
            w, h = im.size
        width = 6.2 if key not in ("gh_approve",) else 4.0
        # keep tall figures on one page
        max_h = 7.6
        height = width * h / w
        if height > max_h:
            width = width * max_h / height
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(width))
        self._place(p._p)
        self.para(caption, 10, True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, after=10)

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            r = p.add_run(line if line else " ")
            set_font(r, 8.5, name="Courier New")
            spacing(p, 1.0, 0, 0, WD_ALIGN_PARAGRAPH.LEFT)
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            ppr.append(shd)
            self._place(p._p)

    def render(self, blocks):
        for b in blocks:
            kind = b[0]
            if kind == "chapter":
                self.chapter(b[1], b[2])
            elif kind == "h2":
                self.h2(b[1])
            elif kind == "h3":
                self.h3(b[1])
            elif kind == "p":
                self.para(b[1])
            elif kind == "bul":
                self.bullets(b[1])
            elif kind == "num":
                self.bullets(b[1], numbered=True)
            elif kind == "table":
                self.table(b[1], b[2], b[3], b[4])
            elif kind == "fig":
                self.figure(b[1], b[2])
            elif kind == "code":
                self.code(b[1])


# ----------------------------------------------------------------------
# appendix code excerpts, taken straight from the repository
# ----------------------------------------------------------------------
def excerpt(path: str, start: str, end: str | None = None, max_lines: int = 60) -> str:
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    s = next(i for i, l in enumerate(lines) if l.startswith(start))
    e = len(lines)
    if end:
        e = next((i for i in range(s + 1, len(lines)) if lines[i].startswith(end)), len(lines))
    return "\n".join(lines[s:e][:max_lines])


# ----------------------------------------------------------------------
# build
# ----------------------------------------------------------------------
def build(toc_pages: dict | None = None) -> Document:
    doc = Document(str(TEMPLATE))
    P = list(doc.paragraphs)       # snapshot: indices refer to the supplied template
    tables = list(doc.tables)
    images = image_paths()

    # ---- cover ------------------------------------------------------
    set_text(P[2], T.TITLE)
    remove(P[3])
    set_text(P[12], T.PROGRAMME)
    for r in P[12].runs:
        r.font.size = Pt(15)
    remove(P[16])
    for row, (reg, name) in zip(tables[0].rows, T.TEAM):
        set_text(row.cells[0].paragraphs[0], reg)
        set_text(row.cells[1].paragraphs[0], name)
    set_text(tables[1].rows[0].cells[0].paragraphs[0], T.GUIDE[0])
    set_text(tables[1].rows[1].cells[0].paragraphs[0], T.GUIDE[1])

    # ---- declaration (written for Gaurish Todi) -----------------------
    set_text(P[27],
             "I hereby declare that the project entitled {} submitted by me, for the award of "
             "the degree of Bachelor of Technology in Computer Science and Engineering with "
             "specialization in Information Security to VIT is a record of bonafide work carried "
             "out by me under the supervision of {}.".format(T.TITLE_PLAIN, T.GUIDE[0]))
    remove(P[43])
    remove(P[94])  # "<Contents, Times New Roman 12 ...>" placeholder under the TOC heading
    clone_after(P[37], "Gaurish Todi (23BCI0262)")

    # ---- certificate -------------------------------------------------
    set_text(P[47],
             "This is to certify that the project entitled {} submitted by Gaurish Todi "
             "(23BCI0262), School of Computer Science and Engineering, VIT, for the award of the "
             "degree of Bachelor of Technology in Computer Science and Engineering with "
             "specialization in Information Security, is a record of bonafide work carried out "
             "under my supervision during Fall Semester 2026-2027, as per the VIT code of "
             "academic and research ethics.".format(T.TITLE_PLAIN))
    remove(P[69])
    spare = [p for p in P[59:66] if not p.text.strip() and "sectPr" not in p._p.xml
             and "w:br" not in p._p.xml]
    for p in spare[:3]:
        remove(p)
    clone_after(P[58], "{}, {}".format(T.GUIDE[0], T.GUIDE[1]))
    set_text(P[73], T.HOD)
    set_text(P[74], "Head of the Department, " + T.HOD_DEPT)

    # ---- acknowledgements (first person, Gaurish Todi) -----------------
    set_text(P[79],
             "I would like to begin by thanking Vellore Institute of Technology (VIT) for the "
             "chance to work on a project like this and for the labs, network and cloud "
             "credits that made real experiments possible. Being able to deploy to an actual "
             "Azure subscription, even a tiny one, changed this from a paper design into "
             "something that had to work.")
    set_text(P[81],
             "My sincere thanks to Dr. Jaisankar N, the Dean of the School of Computer Science "
             "and Engineering (SCOPE), for his support and for the environment that SCOPE "
             "provides to its students.")
    set_text(P[83],
             "I am grateful to {}, Head of the Department of {}, for his guidance and for the "
             "encouragement he has given our programme.".format(T.HOD, T.HOD_DEPT))
    set_text(P[85],
             "I am especially thankful to my project guide, {}, {}, for his patient mentorship. "
             "His questions during the reviews pushed us to measure what we had only claimed, "
             "and to check our framework knowledge against the original documentation, which "
             "made the final work more honest and more useful.".format(T.GUIDE[0], T.GUIDE[1]))
    set_text(P[87],
             "I thank my teammates, Bhumika Singh and Akshat Gupta, for the CI/CD pipeline and "
             "the rule engine this project stands on, and my family and friends for their "
             "patience during the long evenings of debugging. Thank you all for your "
             "contributions and support.")
    set_text(P[91], "Gaurish Todi")

    # ---- table of contents -------------------------------------------
    toc_rows = [("", "Abstract", "abstract")]
    h_by_chapter = {}
    # headings are known only after rendering; use the fixed outline from the template
    outline = [
        ("1.", "INTRODUCTION"), ("", "1.1 Background"), ("", "1.2 Motivations"),
        ("", "1.3 Scope of the Project"),
        ("2.", "PROJECT DESCRIPTION AND GOALS"), ("", "2.1 Literature Review"),
        ("", "2.2 Research Gap"), ("", "2.3 Objectives"), ("", "2.4 Problem Statement"),
        ("", "2.5 Project Plan"),
        ("3.", "TECHNICAL SPECIFICATION"), ("", "3.1 Requirements"), ("", "3.1.1 Functional"),
        ("", "3.1.2 Non-Functional"), ("", "3.2 Feasibility Study"),
        ("", "3.2.1 Technical Feasibility"), ("", "3.2.2 Economic Feasibility"),
        ("", "3.2.3 Social Feasibility"), ("", "3.3 System Specification"),
        ("", "3.3.1 Hardware Specification"), ("", "3.3.2 Software Specification"),
        ("4.", "DESIGN APPROACH AND DETAILS"), ("", "4.1 System Architecture"), ("", "4.2 Design"),
        ("", "4.2.1 Data Flow Diagram"), ("", "4.2.2 Use Case Diagram"),
        ("", "4.2.3 Class Diagram"), ("", "4.2.4 Sequence Diagram"),
        ("5.", "METHODOLOGY AND TESTING"), ("", "5.1 Module Description"), ("", "5.2 Testing"),
        ("6.", "PROJECT DEMONSTRATION"),
        ("7.", "RESULT AND DISCUSSION (COST ANALYSIS AS APPLICABLE)"),
        ("8.", "CONCLUSION"), ("9.", "REFERENCES"), ("", "APPENDIX A – SAMPLE CODE"),
    ]
    entries = [("", "Abstract", "Abstract")] + [(n, t, t) for n, t in outline]
    toc = tables[2]
    while len(toc.rows) - 1 < len(entries):
        toc._tbl.append(copy.deepcopy(toc.rows[-1]._tr))
    while len(toc.rows) - 1 > len(entries):
        toc._tbl.remove(toc.rows[-1]._tr)
    for row, (num, text, key) in zip(toc.rows[1:], entries):
        set_text(row.cells[0].paragraphs[0], num)
        set_text(row.cells[1].paragraphs[0], text)
        page = (toc_pages or {}).get(key, "")
        set_text(row.cells[2].paragraphs[0], page)
        top = (num != "") or key == "Abstract"
        for r in row.cells[1].paragraphs[0].runs:
            r.bold = top

    # ---- list of figures / tables -------------------------------------
    lof = tables[3]
    keep = copy.deepcopy(lof.rows[1]._tr)
    for row in list(lof.rows)[1:]:
        lof._tbl.remove(row._tr)
    for num, title in T.FIGURES:
        lof._tbl.append(copy.deepcopy(keep))
        row = lof.rows[-1]
        set_text(row.cells[0].paragraphs[0], num)
        set_text(row.cells[1].paragraphs[0], title)
        set_text(row.cells[2].paragraphs[0], (toc_pages or {}).get("fig " + num, ""))

    remove(P[105])  # caption-format note
    table_titles = [b for b in T.CH if b[0] == "table"]
    anchor = P[110]
    for i, b in enumerate(table_titles):
        cap = b[1]                       # "Table 2.1 Summary of related work"
        num = cap.split()[1]
        title = cap.split(" ", 2)[2]
        np_ = clone_after(anchor, "{}\t{}\t{}".format(num, title, (toc_pages or {}).get("tab " + num, "")))
        anchor = np_
    remove(P[110])

    # ---- abbreviations ------------------------------------------------
    for p in P[115:129]:
        remove(p)
    prev = P[114]
    set_text(prev, "{}\t{}".format(*T.ABBREVIATIONS[0]))
    for abbr, meaning in T.ABBREVIATIONS[1:]:
        prev = clone_after(prev, "{}\t{}".format(abbr, meaning))

    # ---- symbols and notations ---------------------------------------
    for p in P[133:184]:
        remove(p)
    sym = P[132]
    for r in sym.runs:
        r.font.name = FONT
    first = True
    prev = sym
    for s, meaning in T.SYMBOLS:
        text = "{}\t{}".format(s, meaning)
        if first:
            set_text(sym, text)
            first = False
        else:
            prev = clone_after(prev, text)
    for sp in [sym] + [q for q in P[132]._p.itersiblings()][:0]:
        pass
    from docx.text.paragraph import Paragraph
    node = sym._p
    for _ in T.SYMBOLS:
        para = Paragraph(node, sym._parent)
        for r in para.runs:
            set_font(r, 12)
        para.paragraph_format.left_indent = Inches(1.6)
        para.paragraph_format.first_line_indent = Inches(-1.6)
        para.paragraph_format.tab_stops.add_tab_stop(Inches(1.6))
        node = node.getnext()
    P[184].paragraph_format.page_break_before = True
    # empty filler paragraphs between the abstract and the chapter section
    # break would otherwise spill onto a blank page
    for p in P[189:225]:
        if not p.text.strip() and "sectPr" not in p._p.xml and "w:br" not in p._p.xml:
            remove(p)

    # ---- abstract -----------------------------------------------------
    remove(P[186])
    remove(P[188])
    set_text(P[187], T.ABSTRACT[0])
    spacing(P[187], 1.15, 0, 6, WD_ALIGN_PARAGRAPH.JUSTIFY)
    prev = P[187]
    for para in T.ABSTRACT[1:]:
        prev = clone_after(prev, para)
    kw = clone_after(prev, "Keywords: " + T.KEYWORDS)
    for r in kw.runs:
        r.italic = True

    # ---- chapters (inserted before the section-break paragraph) -----------
    for p in P[226:230]:
        remove(p)
    writer = Writer(doc, P[230]._p, images)
    writer.render(T.CH)

    # ---- references + appendix ---------------------------------------
    set_text(P[231], "CHAPTER 9")
    set_text(P[232], "9. REFERENCES")
    for p in P[233:254]:
        remove(p)
    ref_writer = Writer(doc, P[254]._p, images)
    for i, ref in enumerate(T.REFERENCES, 1):
        p = doc.add_paragraph()
        r = p.add_run("[{}]\t{}".format(i, ref))
        set_font(r, 12)
        spacing(p, 1.15, 0, 5, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
        ref_writer._place(p._p)

    app_writer = Writer(doc, doc.element.body[-1], images)   # before final sectPr
    app_writer.para(T.APPENDIX_INTRO, after=6)
    app_writer.para("Listing A.1 Orchestration core (agent/engine.py)", 10, True,
                    align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=4, after=2, keep_next=True)
    app_writer.code(excerpt("agent/engine.py", "def evaluate_plan", "def evaluate_file", 62))
    app_writer.para("Listing A.2 Wildcard grant rule after the audit fix (agent/rules/caf.py)", 10,
                    True, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=10, after=2,
                    keep_next=True)
    app_writer.code(excerpt("agent/rules/caf.py", "    def _allows_all_actions", "    def check", 30))
    app_writer.para("Listing A.3 Verification of LLM output (agent/reasoner/llm.py)", 10, True,
                    align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=10, after=2, keep_next=True)
    app_writer.code(excerpt("agent/reasoner/llm.py", "        valid_addresses", "    def stats", 30))
    return doc


# ----------------------------------------------------------------------
# page numbers: Word -> PDF -> read back
# ----------------------------------------------------------------------
def to_pdf(docx_path: Path) -> Path:
    pdf = docx_path.with_suffix(".pdf")
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        "$d = $w.Documents.Open('{}'); $d.Fields.Update() | Out-Null; "
        "$d.SaveAs2('{}', 17); $d.Close($false); $w.Quit()"
    ).format(str(docx_path).replace("'", "''"), str(pdf).replace("'", "''"))
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=240)
    return pdf


def read_pages(pdf: Path) -> dict:
    import pymupdf as fitz

    doc = fitz.open(str(pdf))
    texts = []
    shown = []
    for page in doc:
        lines = [l.strip() for l in page.get_text().splitlines() if l.strip()]
        texts.append(lines)
        head = lines[0] if lines else ""
        shown.append(head if len(head) <= 6 else "")
    # first page whose lines contain the chapter-1 heading after the contents
    start = next(i for i, ls in enumerate(texts) if "CHAPTER 1" in ls)
    result = {}
    abs_idx = next(i for i, ls in enumerate(texts) if "ABSTRACT" in ls and i > 3)
    result["Abstract"] = shown[abs_idx] or "i"

    def find(label_variants, from_page=start):
        for i in range(from_page, len(texts)):
            for ln in texts[i]:
                if any(ln.startswith(v) for v in label_variants):
                    return shown[i]
        return ""

    outline_keys = {}
    for b in T.CH:
        if b[0] == "chapter":
            outline_keys[b[2]] = ["{}. {}".format(b[1].split()[-1], b[2]), b[2]]
        elif b[0] in ("h2", "h3"):
            outline_keys[b[1]] = [b[1]]
    outline_keys["REFERENCES"] = ["9. REFERENCES"]
    outline_keys["APPENDIX A – SAMPLE CODE"] = ["APPENDIX A"]
    for key, variants in outline_keys.items():
        result[key] = find(variants)
    # figure and table captions
    for b in T.CH:
        if b[0] == "fig":
            result["fig " + b[2].split()[1]] = find([b[2][:45]])
        elif b[0] == "table":
            result["tab " + b[1].split()[1]] = find([b[1][:45]])
    return result


def main() -> int:
    doc = build()
    doc.save(str(OUTPUT))
    print("pass 1 saved; measuring page numbers with Word ...")
    pdf = to_pdf(OUTPUT)
    pages = read_pages(pdf)
    # chapter keys in the TOC use the bare title text
    doc = build(pages)
    doc.save(str(OUTPUT))
    to_pdf(OUTPUT)
    print("saved", OUTPUT.name)
    missing = [k for k, v in pages.items() if not v]
    if missing:
        print("page number not found for:", missing)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
